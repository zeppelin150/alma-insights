"""Accuracy audit of the in-app Help Center — section ``create``.

Every test here settles a falsifiable claim made by one of the seven articles
under ``assets/help/create/``:

  * ``powerpoint.md``        — PowerPoint: modelling a deck
  * ``deck-outline.md``      — Editing the outline and exporting
  * ``zendesk.md``           — Zendesk: syncing and drafting articles
  * ``zendesk-macros.md``    — Editing macros
  * ``content-studio.md``    — The Content Studio
  * ``studio-no-screen.md``  — The Content Studio has no screen yet
  * ``deferred-previews.md`` — What is deferred

Each test's docstring names the article and quotes (or closely paraphrases)
the claim it settles. Tests marked ``xfail(strict=True)`` assert the ARTICLE's
claim while the code does something else — they are the CONTRADICTED findings
and must stay red-but-tracked until either the article or the code changes.

Headless: no network, no credentials, no QtWebEngine. Qt widgets are exercised
through their public APIs only (never screenshots — offscreen grabs are blank).
"""

from __future__ import annotations

import ast
import json
import os
import re
from pathlib import Path

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from src.data import artifact_store, mermaid_lint, pptx_store, quiz_artifacts  # noqa: E402
from src.data import zendesk_store  # noqa: E402
from src.data.chat_tools import artifact_tools  # noqa: E402

_PROJECT_ROOT = Path(__file__).resolve().parent.parent
_HELP_DIR = _PROJECT_ROOT / "assets" / "help" / "create"
_PAGE_PY = _PROJECT_ROOT / "src" / "ui" / "pages" / "enablement" / "page.py"


# ══════════════════════════════════════════════════════════════════════
#  fixtures
# ══════════════════════════════════════════════════════════════════════

@pytest.fixture(scope="module")
def qapp():
    """One QApplication for the whole module (offscreen)."""
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])
    yield app


@pytest.fixture
def conn(empty_db):
    """Raw connection off the standard ``empty_db`` fixture."""
    return empty_db.conn


class _FakePptxTab:
    """Records what the page handlers push into the PowerPoint view."""

    def __init__(self):
        self.status = []
        self.decks = None
        self.shown = None

    def set_status(self, text):
        self.status.append(text)

    def set_decks(self, decks):
        self.decks = decks

    def show_deck(self, deck):
        self.shown = deck


class _FakeZendeskTab:
    def __init__(self):
        self.status = []
        self.articles = None
        self.macros = None
        self.shown_article = None
        self.shown_macro = None

    def set_status(self, text):
        self.status.append(text)

    def set_articles(self, count, drafts):
        self.articles = (count, drafts)

    def set_macros(self, count, drafts):
        self.macros = (count, drafts)

    def show_article_draft(self, d):
        self.shown_article = d

    def show_macro_draft(self, d):
        self.shown_macro = d


class _FakePage:
    """Minimal stand-in for EnablementPage so its handlers can be called
    unbound — building the real page would start threads and touch settings.

    The two list-refresh helpers delegate to the REAL page implementations so
    the "a new deck appears in the list" style of claim is genuinely exercised.
    """

    def __init__(self, conn, *, demo=False):
        self._c = conn
        self.demo = demo
        self.pptx = _FakePptxTab()
        self.zendesk = _FakeZendeskTab()
        self._zd_client = None

    def _conn(self):
        return self._c

    def _zendesk_client(self):
        return self._zd_client

    def _load_pptx(self):
        return _page_cls()._load_pptx(self)

    def _load_zendesk(self):
        return _page_cls()._load_zendesk(self)

    def _on_pptx_deck_selected(self, deck_id):
        return _page_cls()._on_pptx_deck_selected(self, deck_id)


def _page_cls():
    from src.ui.pages.enablement.page import EnablementPage
    return EnablementPage


def _mk_deck(conn, title, slides):
    return pptx_store.save_deck(conn, title=title,
                                outline={"title": title, "slides": slides})


# ══════════════════════════════════════════════════════════════════════
#  article: powerpoint.md — "PowerPoint: modelling a deck"
# ══════════════════════════════════════════════════════════════════════

def test_pptx_tab_offers_both_a_topic_route_and_a_document_route(qapp):
    """powerpoint.md: "There are two routes into a new deck" — a topic field
    with a Model deck button, and dropping a document on the tab."""
    from src.ui.pages.enablement.pptx_tab import PptxPage
    page = PptxPage()
    # topic route: the line edit and the button both fire model_topic_requested
    seen = []
    page.model_topic_requested.connect(seen.append)
    page._topic.setText("Aetna eligibility rechecks")
    page._on_model_topic()
    assert seen == ["Aetna eligibility rechecks"]
    # document route: the tab accepts drops and relays the local path
    assert page.acceptDrops() is True
    dropped = []
    page.doc_dropped.connect(dropped.append)
    assert hasattr(page, "dropEvent")


def test_topic_route_calls_the_language_model_exactly_once(conn):
    """powerpoint.md: modelling "from a topic" "calls the language model, so it
    needs a working model connection"."""
    calls = []

    class _LLM:
        def generate(self, prompt):
            calls.append(prompt)
            return json.dumps({"title": "Aetna 101",
                               "slides": [{"title": "Why", "bullets": ["a"]}]})

    res = pptx_store.generate_deck_from_topic(conn, "Aetna rechecks", _LLM())
    assert res["ok"] is True
    assert len(calls) == 1, "topic route must go through the LLM"


def test_topic_route_reports_a_missing_model_client_on_the_status_line(conn):
    """powerpoint.md ("If it doesn't"): "Modelling from a topic fails and
    mentions a missing model client"."""
    page = _FakePage(conn)
    _page_cls()._on_pptx_modeled(page, {"ok": False, "error": "no_llm_client"})
    assert page.pptx.status[-1] == "Modeling failed: no_llm_client"


def test_document_route_uses_no_language_model_at_all(conn, tmp_path, monkeypatch):
    """powerpoint.md: "This route uses no language model at all." Proven by
    making every LLM client build explode and dropping a doc anyway."""
    from src.gemini import client_factory

    def _boom(*a, **k):
        raise AssertionError("the document route must not build an LLM client")

    monkeypatch.setattr(client_factory, "build_client_for_task", _boom)
    doc = tmp_path / "refunds_policy.md"
    doc.write_text("# Refund windows\n\n- 30 days\n- No exceptions\n",
                   encoding="utf-8")

    page = _FakePage(conn)
    _page_cls()._on_pptx_doc_dropped(page, str(doc))

    decks = pptx_store.list_decks(conn)
    assert len(decks) == 1
    assert decks[0]["source_ref"] == "upload:refunds_policy.md"


def test_document_route_turns_every_heading_into_its_own_slide():
    """powerpoint.md table: "Any heading → A new slide, titled with that
    heading"."""
    md = "# Alpha\n\ntext a\n\n## Beta\n\ntext b\n\n### Gamma\n\ntext c\n"
    out = pptx_store.outline_from_markdown("Doc.md", md)
    titles = [s["title"] for s in out["slides"]]
    assert titles == ["Doc", "Alpha", "Beta", "Gamma"]


def test_document_route_turns_bullets_and_numbered_items_into_bullets():
    """powerpoint.md table: "A bullet or numbered item → A bullet on the
    current slide"."""
    md = "# Steps\n- first\n* second\n+ third\n1. fourth\n2) fifth\n"
    out = pptx_store.outline_from_markdown("Doc.md", md)
    steps = [s for s in out["slides"] if s["title"] == "Steps"][0]
    assert steps["bullets"] == ["first", "second", "third", "fourth", "fifth"]


def test_document_route_turns_a_paragraph_into_a_bullet():
    """powerpoint.md table: "A short paragraph → A bullet on the current
    slide"."""
    md = "# Context\n\nEligibility must be rechecked every 30 days.\n"
    out = pptx_store.outline_from_markdown("Doc.md", md)
    ctx = [s for s in out["slides"] if s["title"] == "Context"][0]
    assert ctx["bullets"] == ["Eligibility must be rechecked every 30 days."]


def test_document_route_skips_tables_and_raw_markup():
    """powerpoint.md table: "Tables and raw markup → Skipped"."""
    md = ("# Data\n"
          "| col | col |\n"
          "|---|---|\n"
          "<div class='x'>raw</div>\n"
          "kept line\n")
    out = pptx_store.outline_from_markdown("Doc.md", md)
    data = [s for s in out["slides"] if s["title"] == "Data"][0]
    assert data["bullets"] == ["kept line"]


def test_document_route_caps_each_slide_at_eight_bullets():
    """powerpoint.md: "The document route caps each slide at eight bullets"."""
    md = "# Many\n" + "".join(f"- item {i}\n" for i in range(20))
    out = pptx_store.outline_from_markdown("Doc.md", md)
    many = [s for s in out["slides"] if s["title"] == "Many"][0]
    assert len(many["bullets"]) == 8
    assert many["bullets"][-1] == "item 7"


def test_document_route_trims_a_bullet_to_160_characters():
    """powerpoint.md: "trims any single bullet to roughly 160 characters"."""
    long = "x" * 400
    out = pptx_store.outline_from_markdown("Doc.md", f"# Long\n- {long}\n")
    lng = [s for s in out["slides"] if s["title"] == "Long"][0]
    assert len(lng["bullets"][0]) == 160


def test_document_route_puts_a_cover_slide_with_the_document_name_first():
    """powerpoint.md: "A cover slide carrying the document's name goes
    first". The extension and separators are cleaned off the file name."""
    out = pptx_store.outline_from_markdown("refunds_policy_v2.docx",
                                           "# Alpha\n- one\n")
    assert out["slides"][0]["title"] == "refunds policy v2"
    assert out["title"] == "refunds policy v2"


def test_document_route_drops_empty_slides_but_keeps_the_cover():
    """powerpoint.md: "empty slides are dropped"."""
    md = "# Has content\n- one\n\n# Empty section\n\n# Also empty\n"
    out = pptx_store.outline_from_markdown("Doc.md", md)
    titles = [s["title"] for s in out["slides"]]
    assert titles == ["Doc", "Has content"]
    assert out["slides"][0]["bullets"] == []   # the cover survives empty


def test_deck_list_shows_every_deck_with_its_slide_count(qapp):
    """powerpoint.md: "The list on the left shows every deck with its slide
    count"."""
    from src.ui.pages.enablement.pptx_tab import PptxPage
    page = PptxPage()
    page.set_decks([{"id": 1, "title": "Aetna 101", "slide_count": 4},
                    {"id": 2, "title": "Refunds", "slide_count": 9}])
    labels = [page._list.item(i).text() for i in range(page._list.count())]
    assert labels == ["Aetna 101  ·  4 slides", "Refunds  ·  9 slides"]


def test_dropping_an_unreadable_path_reports_it_could_not_be_read(conn):
    """powerpoint.md ("If it doesn't"): "Dragging a file in reports it could
    not be read"."""
    page = _FakePage(conn)
    _page_cls()._on_pptx_doc_dropped(page, "C:/definitely/not/a/file.md")
    assert page.pptx.status == ["Could not read that file."]
    assert pptx_store.list_decks(conn) == []


def test_a_new_deck_appears_in_the_list_and_opens_with_a_naming_status(
        conn, tmp_path):
    """powerpoint.md: "A new deck should appear in the list and open
    immediately, with the status line naming it and its slide count.\""""
    doc = tmp_path / "Onboarding.md"
    doc.write_text("# Day one\n- badge\n\n# Day two\n- systems\n",
                   encoding="utf-8")
    page = _FakePage(conn)
    _page_cls()._on_pptx_doc_dropped(page, str(doc))

    assert page.pptx.decks is not None and len(page.pptx.decks) == 1   # listed
    assert page.pptx.shown is not None                                  # opened
    status = page.pptx.status[-1]
    assert "Onboarding" in status and "3 slides" in status


def test_dropping_an_unsupported_format_reports_it_could_not_be_modeled(
        conn, tmp_path):
    """powerpoint.md ("If it doesn't"): an unsupported format now reports it
    couldn't model a deck rather than producing a nonsense deck. The
    deck-modelling path reads with ``strict=True``, so ``read_document`` raises
    ``UnsupportedDocumentError`` for a binary/unreadable format — caught either
    by its extension (``_BINARY_EXT``) or by a NUL-byte content sniff
    (doc_reader.py:100-135). The page handler's ``except`` then surfaces
    "Couldn't model a deck: …" and writes nothing (page.py:_on_pptx_doc_dropped)."""
    from src.data import doc_reader

    # A tiny binary file named .pdf, with a NUL byte, is refused in strict mode.
    pdf = tmp_path / "policy.pdf"
    pdf.write_bytes(b"%PDF-1.4\n\x00\x01 not real text")
    with pytest.raises(doc_reader.UnsupportedDocumentError):
        doc_reader.read_document(str(pdf), strict=True)

    # End to end: dropping it on the tab reports a modelling failure — the old
    # silent "success" of a deck full of decoded bytes is gone.
    page = _FakePage(conn)
    _page_cls()._on_pptx_doc_dropped(page, str(pdf))

    assert any("Couldn't model a deck" in s for s in page.pptx.status), \
        f"expected a modelling-failure status; got {page.pptx.status}"
    assert not any("Modeled" in s for s in page.pptx.status), \
        f"a success status was reported unexpectedly: {page.pptx.status}"

    # No deck was written from the unreadable bytes.
    assert pptx_store.list_decks(conn) == []


def test_a_freshly_modelled_deck_is_saved_as_a_draft_not_exported(conn, tmp_path):
    """powerpoint.md: "It is saved as a draft at that point; nothing has been
    exported.\""""
    doc = tmp_path / "Doc.md"
    doc.write_text("# A\n- one\n", encoding="utf-8")
    page = _FakePage(conn)
    _page_cls()._on_pptx_doc_dropped(page, str(doc))
    deck = pptx_store.list_decks(conn)[0]
    assert deck["status"] != "exported"
    assert not deck["file_path"]


# ══════════════════════════════════════════════════════════════════════
#  article: deck-outline.md — "Editing the outline and exporting"
# ══════════════════════════════════════════════════════════════════════

def test_outline_hash_prefix_opens_a_new_slide():
    """deck-outline.md: "A line starting with `# ` opens a new slide.\""""
    from src.ui.pages.enablement.pptx_tab import text_to_outline
    out = text_to_outline("Deck", "# One\n# Two\n# Three\n")
    assert [s["title"] for s in out["slides"]] == ["One", "Two", "Three"]


def test_outline_dash_prefix_adds_a_bullet_to_the_open_slide():
    """deck-outline.md: "A line starting with `- ` adds a bullet to whichever
    slide is open.\""""
    from src.ui.pages.enablement.pptx_tab import text_to_outline
    out = text_to_outline("Deck", "# One\n- a\n- b\n# Two\n- c\n")
    assert out["slides"][0]["bullets"] == ["a", "b"]
    assert out["slides"][1]["bullets"] == ["c"]


def test_outline_ignores_every_other_line():
    """deck-outline.md: "**Every other line is ignored.** Blank lines, notes to
    yourself, prose paragraphs — none of it survives a save.\""""
    from src.ui.pages.enablement.pptx_tab import text_to_outline
    text = ("# Slide\n"
            "- kept\n"
            "\n"
            "a prose paragraph\n"
            "  - indented bullet\n"
            "* asterisk bullet\n"
            "## h2 heading\n"
            "#no space\n")
    out = pptx_store._normalize_outline(text_to_outline("Deck", text))
    assert len(out["slides"]) == 1
    assert out["slides"][0]["bullets"] == ["kept"]


def test_outline_bullet_before_any_title_gets_a_placeholder_slide():
    """deck-outline.md: "A bullet written before any slide title creates a
    slide of its own with a placeholder title.\""""
    from src.ui.pages.enablement.pptx_tab import text_to_outline
    out = text_to_outline("Deck", "- orphan\n# Real\n- b\n")
    assert out["slides"][0]["title"] == "Slide"
    assert out["slides"][0]["bullets"] == ["orphan"]


def test_slide_view_rerenders_from_unsaved_editor_text(qapp):
    """deck-outline.md: "Switching to the slide view re-renders from whatever
    is currently in the editor, including edits you have not saved.\""""
    from src.ui.pages.enablement.pptx_tab import PptxPage
    page = PptxPage()
    page.show_deck({"id": 1, "title": "D",
                    "outline": {"title": "D", "slides": [{"title": "A",
                                                          "bullets": []}]}})
    page._set_view("outline")
    page._outline.setPlainText("# A\n# B\n# C\n")     # unsaved typing
    page._set_view("slides")
    # render_slides adds one card per slide plus a trailing stretch item
    assert page._slides_layout.count() == 4
    assert page._stack.currentWidget() is page._scroll


def test_saving_the_outline_persists_it_and_reports_the_stored_count(conn):
    """deck-outline.md: "Saving is what persists it, and the status line
    confirms with the slide count it stored.\""""
    deck_id = _mk_deck(conn, "D", [{"title": "A", "bullets": ["x"]}])
    page = _FakePage(conn)
    new_outline = {"title": "D2", "slides": [{"title": "A", "bullets": ["x"]},
                                             {"title": "B", "bullets": ["y"]}]}
    _page_cls()._on_pptx_outline_saved(page, deck_id, "D2", new_outline)

    assert page.pptx.status[-1] == "Saved — 2 slides."
    stored = pptx_store.get_deck(conn, deck_id)
    assert stored["title"] == "D2"
    assert len(stored["outline"]["slides"]) == 2


def test_saving_never_changes_bullet_wording_or_order(conn):
    """deck-outline.md: "Saving should never change your bullets' wording or
    order.\""""
    bullets = ["Zulu first", "alpha SECOND", "  spaced  third  "]
    deck_id = _mk_deck(conn, "D", [{"title": "S", "bullets": list(bullets)}])
    pptx_store.update_deck_outline(
        conn, deck_id, title="D",
        outline={"title": "D", "slides": [{"title": "S",
                                           "bullets": list(bullets)}]})
    stored = pptx_store.get_deck(conn, deck_id)
    assert stored["outline"]["slides"][0]["bullets"] == bullets


def test_the_focus_overlay_commits_through_the_same_save_signal(qapp):
    """deck-outline.md: "There is also a focus view that opens the deck in an
    overlay ...; committing there saves the same way.\""""
    from src.ui.pages.enablement.pptx_tab import PptxPage
    page = PptxPage()
    saved = []
    page.outline_saved.connect(lambda i, t, o: saved.append((i, t, o)))
    page._on_expand_committed(7, "Focused", {"title": "Focused",
                                             "slides": [{"title": "A",
                                                         "bullets": ["b"]}]})
    assert len(saved) == 1
    assert saved[0][0] == 7 and saved[0][1] == "Focused"


def test_export_asks_where_to_put_the_file_and_cancelling_writes_nothing(
        conn, monkeypatch):
    """deck-outline.md: "Exporting asks you where to put the file". Cancelling
    the dialog must leave the deck untouched."""
    from PySide6.QtWidgets import QFileDialog
    deck_id = _mk_deck(conn, "D", [{"title": "A", "bullets": ["x"]}])
    asked = []

    def _cancelled(parent, caption, default, filt):
        asked.append((caption, default, filt))
        return "", ""

    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(_cancelled))
    page = _FakePage(conn)
    _page_cls()._on_pptx_export(page, deck_id)

    assert asked and asked[0][0] == "Export deck"
    assert asked[0][1] == "D.pptx"
    assert pptx_store.get_deck(conn, deck_id)["status"] != "exported"
    assert page.pptx.status == []


def test_export_writes_a_real_pptx_and_records_the_path(conn, tmp_path,
                                                        monkeypatch):
    """deck-outline.md: "writes a real .pptx" and "Exporting marks the deck
    exported and records the path.\""""
    from pptx import Presentation
    from PySide6.QtWidgets import QFileDialog
    deck_id = _mk_deck(conn, "Aetna 101",
                       [{"title": "A", "bullets": ["x"]},
                        {"title": "B", "bullets": ["y"]}])
    out = tmp_path / "aetna.pptx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a: (str(out), "")))
    page = _FakePage(conn)
    _page_cls()._on_pptx_export(page, deck_id)

    assert out.exists()
    assert len(Presentation(str(out)).slides) == 3      # 2 + cover
    stored = pptx_store.get_deck(conn, deck_id)
    assert stored["status"] == "exported"
    assert stored["file_path"] == str(out)
    assert page.pptx.status[-1].startswith("Exported 3 slides → ")


def test_export_adds_a_cover_slide_carrying_the_deck_title(conn, tmp_path):
    """deck-outline.md: "A **cover slide carrying the deck title is added at
    the top**".
    """
    from pptx import Presentation
    deck_id = _mk_deck(conn, "Refund Policy",
                       [{"title": "Windows", "bullets": ["30 days"]}])
    out = tmp_path / "d.pptx"
    pptx_store.export_pptx(conn, deck_id, str(out))
    prs = Presentation(str(out))
    assert prs.slides[0].shapes.title.text == "Refund Policy"
    assert prs.slides[1].shapes.title.text == "Windows"


def test_export_count_is_exactly_one_more_than_the_list_count(conn, tmp_path):
    """deck-outline.md: "The count difference between the list and the export
    message should be exactly one, and always in the same direction.\""""
    for n in (1, 3, 7):
        deck_id = _mk_deck(conn, f"D{n}",
                           [{"title": f"S{i}", "bullets": ["b"]}
                            for i in range(n)])
        listed = [d for d in pptx_store.list_decks(conn)
                  if d["id"] == deck_id][0]["slide_count"]
        res = pptx_store.export_pptx(conn, deck_id, str(tmp_path / f"{n}.pptx"))
        assert listed == n
        assert res["slides"] - listed == 1


def test_export_uses_the_template_file_when_one_is_present(conn, tmp_path):
    """deck-outline.md: "A branded template is used when one is present." The
    template genuinely drives the exported file — proven by renaming a layout
    in a copy of the shipped template and finding that name in the output."""
    import shutil
    from pptx import Presentation
    assert pptx_store._DECK_TEMPLATE.exists(), "no template is shipped"

    marked = tmp_path / "marked.pptx"
    shutil.copyfile(pptx_store._DECK_TEMPLATE, marked)
    prs = Presentation(str(marked))
    prs.slide_layouts[0].name = "ALMA-BRAND-MARKER"
    prs.save(str(marked))

    deck_id = _mk_deck(conn, "D", [{"title": "A", "bullets": ["x"]}])
    out = tmp_path / "d.pptx"
    pptx_store.export_pptx(conn, deck_id, str(out), template_path=str(marked))
    names = [layout.name for layout in Presentation(str(out)).slide_layouts]
    assert "ALMA-BRAND-MARKER" in names


def test_the_shipped_brand_template_is_the_python_pptx_default_placeholder(
        conn, tmp_path):
    """deck-outline.md ("If it doesn't"): the CORRECTED claim — the shipped
    ``assets/templates/renn_deck.pptx`` is python-pptx's own default deck
    re-saved (a placeholder, as pptx_store.py:26 documents). Its theme1.xml
    hashes identically to a fresh default deck, its eleven layout names match,
    and its slide size matches. So an export made WITH the template is
    indistinguishable from the default-layout fallback, and a deck that "looks
    unbranded" does NOT imply the template was missing — the article now says
    decks are unbranded because the template is a placeholder, not a failure."""
    import hashlib
    import zipfile
    from pptx import Presentation

    def _theme_hash(path):
        with zipfile.ZipFile(path) as z:
            name = next(n for n in z.namelist() if n.endswith("theme1.xml"))
            return hashlib.sha256(z.read(name)).hexdigest()

    plain = tmp_path / "plain.pptx"
    Presentation().save(str(plain))

    # Same theme XML, byte-for-byte.
    assert _theme_hash(pptx_store._DECK_TEMPLATE) == _theme_hash(str(plain))

    shipped = Presentation(str(pptx_store._DECK_TEMPLATE))
    default = Presentation(str(plain))
    # Same layout names and same slide size — nothing branded was added.
    assert ([l.name for l in shipped.slide_layouts]
            == [l.name for l in default.slide_layouts])
    assert (shipped.slide_width, shipped.slide_height) == \
        (default.slide_width, default.slide_height)


def test_export_falls_back_to_a_default_layout_when_the_template_is_corrupt(
        conn, tmp_path):
    """deck-outline.md: "if the template file is missing or unreadable, export
    quietly falls back to a plain default layout rather than failing.\""""
    corrupt = tmp_path / "broken.pptx"
    corrupt.write_bytes(b"this is not a pptx at all")
    deck_id = _mk_deck(conn, "D", [{"title": "A", "bullets": ["x"]}])
    out = tmp_path / "d.pptx"
    res = pptx_store.export_pptx(conn, deck_id, str(out),
                                 template_path=str(corrupt))
    assert res["ok"] is True
    assert out.exists()


def test_export_writes_speaker_notes_when_the_outline_carries_them(conn,
                                                                   tmp_path):
    """deck-outline.md: "Speaker notes are written to the .pptx when the
    outline carries them.\""""
    from pptx import Presentation
    deck_id = _mk_deck(conn, "D", [{"title": "A", "bullets": ["x"],
                                    "notes": "Pause here for questions."}])
    out = tmp_path / "d.pptx"
    pptx_store.export_pptx(conn, deck_id, str(out))
    prs = Presentation(str(out))
    assert "Pause here for questions." in prs.slides[1].notes_slide.notes_text_frame.text


def test_export_does_not_change_the_stored_outline(conn, tmp_path):
    """deck-outline.md: "Exporting should never change the stored outline — it
    only reads it.\""""
    slides = [{"title": "A", "bullets": ["one", "two"], "notes": "n"}]
    deck_id = _mk_deck(conn, "D", slides)
    before = json.loads(conn.execute(
        "SELECT outline_json FROM pptx_decks WHERE id=?", (deck_id,)).fetchone()[0])
    pptx_store.export_pptx(conn, deck_id, str(tmp_path / "d.pptx"))
    after = json.loads(conn.execute(
        "SELECT outline_json FROM pptx_decks WHERE id=?", (deck_id,)).fetchone()[0])
    assert before == after


def test_reopening_an_exported_deck_shows_it_as_exported(qapp):
    """deck-outline.md: "Once exported, the deck's status shows as exported
    when you reopen it.\""""
    from src.ui.pages.enablement.pptx_tab import PptxPage
    page = PptxPage()
    page.show_deck({"id": 1, "title": "D", "status": "exported",
                    "outline": {"title": "D", "slides": []}})
    assert page._status.text() == "Exported."
    page.show_deck({"id": 2, "title": "E", "status": "pending",
                    "outline": {"title": "E", "slides": []}})
    assert page._status.text() == "Draft."


def test_saving_from_the_outline_tab_drops_speaker_notes(conn):
    """deck-outline.md ("If it doesn't"): "Speaker notes vanished after you
    saved. This is a real limitation ... a deck that arrived with speaker notes
    ... loses them the first time you save from this tab.\""""
    from src.ui.pages.enablement.pptx_tab import outline_to_text, text_to_outline
    deck_id = _mk_deck(conn, "D", [{"title": "A", "bullets": ["x"],
                                    "notes": "say this out loud"}])
    deck = pptx_store.get_deck(conn, deck_id)
    assert deck["outline"]["slides"][0]["notes"] == "say this out loud"

    # what the tab does: outline → editor text → outline → save
    round_tripped = text_to_outline(deck["title"],
                                    outline_to_text(deck["outline"]))
    pptx_store.update_deck_outline(conn, deck_id, title=deck["title"],
                                   outline=round_tripped)
    reloaded = pptx_store.get_deck(conn, deck_id)
    assert "notes" not in reloaded["outline"]["slides"][0]


def test_export_without_python_pptx_reports_the_missing_library(conn, tmp_path,
                                                                monkeypatch):
    """deck-outline.md ("If it doesn't"): "Export fails saying a PowerPoint
    library is missing.\""""
    import builtins
    deck_id = _mk_deck(conn, "D", [{"title": "A", "bullets": ["x"]}])
    real_import = builtins.__import__

    def _no_pptx(name, *args, **kwargs):
        if name == "pptx":
            raise ImportError("No module named 'pptx'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", _no_pptx)
    res = pptx_store.export_pptx(conn, deck_id, str(tmp_path / "d.pptx"))
    assert res["ok"] is False
    assert "pptx" in res["error"] and "not installed" in res["error"]


def test_export_of_a_missing_deck_reports_deck_not_found(conn, tmp_path):
    """deck-outline.md ("If it doesn't"): "Export fails saying the deck was not
    found.\""""
    res = pptx_store.export_pptx(conn, 999999, str(tmp_path / "d.pptx"))
    assert res == {"ok": False, "error": "deck_not_found"}


# ══════════════════════════════════════════════════════════════════════
#  article: zendesk.md — "Zendesk: syncing and drafting articles"
# ══════════════════════════════════════════════════════════════════════

class _FakeZendeskClient:
    """Records every call so a test can prove sync never writes."""

    def __init__(self, *, articles=None, macros=None):
        self._articles = articles or []
        self._macros = macros or []
        self.calls = []

    # reads
    def get_articles(self, *, locale="en-us", per_page=100):
        self.calls.append(("get_articles", locale, per_page))
        return self._articles

    def get_articles_paged(self, **kw):
        self.calls.append(("get_articles_paged", kw))
        return self._articles, len(self._articles)

    def list_macros(self, *, per_page=100):
        self.calls.append(("list_macros", per_page))
        return self._macros

    def list_macros_paged(self, **kw):
        self.calls.append(("list_macros_paged", kw))
        return self._macros, len(self._macros)

    # writes
    def create_article(self, section_id, title, body, *, locale="en-us"):
        self.calls.append(("create_article", section_id, title, body))
        return {"id": 4242}

    def update_article(self, article_id, *, title=None, body=None,
                       locale="en-us"):
        self.calls.append(("update_article", article_id, title, body))
        return {"id": article_id}

    def create_macro(self, name, actions, *, description=None):
        self.calls.append(("create_macro", name, actions, description))
        return {"id": 777}

    def update_macro(self, macro_id, *, name=None, actions=None):
        self.calls.append(("update_macro", macro_id, name, actions))
        return {"id": macro_id}


_WRITE_CALLS = {"create_article", "update_article", "create_macro",
                "update_macro"}


def test_sync_reads_and_never_writes(conn):
    """zendesk.md: "**Sync reads. It never writes.**" and "Nothing goes to
    Zendesk during a sync"."""
    client = _FakeZendeskClient(
        articles=[{"id": 1, "title": "A", "body": "<p>b</p>",
                   "section_id": 9, "updated_at": "2026-01-01"}],
        macros=[{"id": 5, "title": "M", "actions": [], "updated_at": "2026-01-01"}])
    zendesk_store.sync_articles(conn, client)
    zendesk_store.sync_macros(conn, client)
    assert not [c for c in client.calls if c[0] in _WRITE_CALLS]


def test_sync_caches_articles_and_macros_and_refreshes_existing_rows(conn):
    """zendesk.md: "Syncing pulls Help Center articles and account macros into
    a local cache inside the app and refreshes rows already there.\""""
    c1 = _FakeZendeskClient(
        articles=[{"id": 1, "title": "Old title", "body": "b",
                   "section_id": 9, "updated_at": "2026-01-01"}],
        macros=[{"id": 5, "title": "Old macro", "actions": [],
                 "updated_at": "2026-01-01"}])
    assert zendesk_store.sync_articles(conn, c1) == {"ok": True, "count": 1}
    assert zendesk_store.sync_macros(conn, c1) == {"ok": True, "count": 1}

    c2 = _FakeZendeskClient(
        articles=[{"id": 1, "title": "New title", "body": "b2",
                   "section_id": 9, "updated_at": "2026-02-01"}],
        macros=[{"id": 5, "title": "New macro", "actions": [],
                 "updated_at": "2026-02-01"}])
    zendesk_store.sync_articles(conn, c2)
    zendesk_store.sync_macros(conn, c2)

    arts = zendesk_store.list_articles(conn)
    macs = zendesk_store.list_macros(conn)
    assert len(arts) == 1 and arts[0]["title"] == "New title"
    assert len(macs) == 1 and macs[0]["name"] == "New macro"


def test_sync_creates_no_drafts(conn):
    """zendesk.md: "no drafts are created from what it finds.\""""
    client = _FakeZendeskClient(
        articles=[{"id": 1, "title": "A", "body": "b", "updated_at": "x"}],
        macros=[{"id": 5, "title": "M", "actions": [], "updated_at": "x"}])
    zendesk_store.sync_articles(conn, client)
    zendesk_store.sync_macros(conn, client)
    assert zendesk_store.list_article_drafts(conn) == []
    assert zendesk_store.list_macro_drafts(conn) == []


def test_sync_status_line_reports_how_many_of_each_it_pulled(conn):
    """zendesk.md: "The status line reports how many of each it pulled.\""""
    page = _FakePage(conn)
    _page_cls()._on_zendesk_synced(page, {"ok": True, "articles": 12,
                                          "macros": 5})
    assert page.zendesk.status[0] == "Synced 12 articles, 5 macros."


def test_sync_takes_only_one_page_from_each_endpoint(conn):
    """zendesk.md: "Sync takes one page from each endpoint, so roughly the
    first hundred articles and hundred macros rather than your whole Help
    Center.\""""
    client = _FakeZendeskClient(articles=[], macros=[])
    zendesk_store.sync_articles(conn, client)
    zendesk_store.sync_macros(conn, client)
    names = [c[0] for c in client.calls]
    assert names == ["get_articles", "list_macros"]
    assert "get_articles_paged" not in names
    assert "list_macros_paged" not in names
    # and the single page requested is the 100-item default
    assert client.calls[0][2] == 100
    assert client.calls[1][1] == 100


def test_article_count_comes_from_the_cache_and_the_list_from_drafts(conn,
                                                                     qapp):
    """zendesk.md: "The count beside the heading comes from the cache; the list
    below it comes from pending drafts", "which is why a sync can report a
    healthy article count and leave the list empty.\""""
    from src.ui.pages.enablement.zendesk_tab import ZendeskPage
    client = _FakeZendeskClient(
        articles=[{"id": i, "title": f"A{i}", "body": "b", "updated_at": "x"}
                  for i in range(1, 4)])
    zendesk_store.sync_articles(conn, client)

    tab = ZendeskPage()
    page = _FakePage(conn)
    page.zendesk = tab
    _page_cls()._load_zendesk(page)

    assert "3 articles synced" in tab._sync_status.text()
    assert tab._a_list.count() == 0          # cache full, draft list empty


def test_selecting_an_article_draft_loads_its_title_and_body(conn, qapp):
    """zendesk.md: "Selecting a draft loads its title and body into the
    editor.\""""
    from src.ui.pages.enablement.zendesk_tab import ZendeskPage
    draft_id = zendesk_store.save_article_draft(
        conn, title="Refund windows", body="# Refunds\n\nThirty days.")
    tab = ZendeskPage()
    page = _FakePage(conn)
    page.zendesk = tab
    _page_cls()._on_zd_article_selected(page, draft_id)

    assert tab._a_title.text() == "Refund windows"
    assert "Refunds" in tab._a_body.to_markdown()


def test_saving_an_article_draft_stores_it_locally(conn):
    """zendesk.md: "Saving stores it locally.\""""
    draft_id = zendesk_store.save_article_draft(conn, title="T", body="B")
    page = _FakePage(conn)
    _page_cls()._on_zd_article_saved(page, draft_id, "T2", "B2")
    d = zendesk_store.get_article_draft(conn, draft_id)
    assert (d["title"], d["body"]) == ("T2", "B2")
    assert d["status"] == "pending"


def test_pushing_converts_the_body_to_html(conn):
    """zendesk.md: "Pushing converts the body to HTML and sends it".\""""
    client = _FakeZendeskClient()
    draft_id = zendesk_store.save_article_draft(
        conn, title="T", body="# Head\n\nSome **bold** text", article_id=17)
    zendesk_store.publish_article_draft(conn, draft_id, zendesk_client=client)
    sent = [c for c in client.calls if c[0] == "update_article"][0]
    assert "<h1>Head</h1>" in sent[3]
    assert "<strong>bold</strong>" in sent[3]
    assert "**bold**" not in sent[3]


def test_a_linked_draft_updates_and_an_unlinked_draft_creates(conn):
    """zendesk.md: "a draft already linked to a live article updates that
    article, and an unlinked draft creates a new one.\""""
    client = _FakeZendeskClient()
    linked = zendesk_store.save_article_draft(conn, title="L", body="b",
                                              article_id=42)
    zendesk_store.publish_article_draft(conn, linked, zendesk_client=client)
    assert [c[0] for c in client.calls] == ["update_article"]
    assert client.calls[0][1] == 42

    client2 = _FakeZendeskClient()
    unlinked = zendesk_store.save_article_draft(conn, title="N", body="b",
                                                section_id=9)
    res = zendesk_store.publish_article_draft(conn, unlinked,
                                              zendesk_client=client2)
    assert [c[0] for c in client2.calls] == ["create_article"]
    assert client2.calls[0][1] == 9
    assert res["article_id"] == 4242      # the new id is stored back


def test_pushed_drafts_leave_the_pending_list(conn):
    """zendesk.md: "Pushed drafts leave the list, because the list only shows
    pending ones.\""""
    draft_id = zendesk_store.save_article_draft(conn, title="T", body="b")
    assert [d["id"] for d in zendesk_store.list_article_drafts(conn)] == [draft_id]
    zendesk_store.publish_article_draft(conn, draft_id, zendesk_client=None)
    assert zendesk_store.list_article_drafts(conn) == []


def test_creating_a_new_article_without_a_section_is_rejected(conn):
    """zendesk.md: "An unlinked draft pushed against a live Zendesk returns an
    error saying a section is required.\""""
    client = _FakeZendeskClient()
    draft_id = zendesk_store.save_article_draft(conn, title="T", body="b")
    res = zendesk_store.publish_article_draft(conn, draft_id,
                                              zendesk_client=client)
    assert res == {"ok": False, "error": "section_id_required_to_create"}
    assert client.calls == []
    assert zendesk_store.get_article_draft(conn, draft_id)["status"] == "pending"


def test_the_zendesk_tab_offers_no_way_to_choose_a_section(qapp):
    """zendesk.md: "**Creating a new article needs a section**, and this tab
    has no way to choose one.\""""
    from src.ui.pages.enablement import zendesk_tab as zt
    from src.ui.pages.enablement.zendesk_tab import ZendeskPage
    from PySide6.QtWidgets import QComboBox, QLineEdit
    tab = ZendeskPage()
    # no combo boxes at all, and no field that asks for a section
    assert tab.findChildren(QComboBox) == []
    hints = " ".join(w.placeholderText()
                     for w in tab.findChildren(QLineEdit)).lower()
    assert "section" not in hints
    # "section" appears in the tab only as the shared ``section_label`` helper,
    # never as a standalone concept the user can set
    tab_src = Path(zt.__file__).read_text(encoding="utf-8").lower()
    assert re.search(r"\bsection\b", tab_src) is None
    assert "section_id" not in tab_src
    # and the page handler pushes without ever supplying one
    push_src = _PAGE_PY.read_text(encoding="utf-8").split(
        "def _on_zd_article_push")[1].split("    def ")[0]
    assert "publish_article_draft" in push_src
    assert "section" not in push_src


def test_pushing_with_no_connection_marks_the_draft_pushed_anyway(conn):
    """zendesk.md: "**Pushing with no Zendesk connection marks the draft pushed
    anyway.** No API call is attempted, the draft is flagged as pushed
    locally.\""""
    draft_id = zendesk_store.save_article_draft(conn, title="T", body="b")
    res = zendesk_store.publish_article_draft(conn, draft_id,
                                              zendesk_client=None)
    assert res["ok"] is True
    assert zendesk_store.get_article_draft(conn, draft_id)["status"] == "pushed"
    assert zendesk_store.get_article_draft(conn, draft_id)["article_id"] is None


def test_a_pushed_article_draft_disables_the_push_button(qapp):
    """zendesk.md: "the button disables. \"Pushed\" is therefore not proof that
    anything reached Zendesk.\""""
    from src.ui.pages.enablement.zendesk_tab import ZendeskPage
    tab = ZendeskPage()
    tab.show_article_draft({"id": 1, "title": "T", "body": "b",
                            "status": "pushed"})
    assert tab._a_push.text() == "Pushed"
    assert tab._a_push.isEnabled() is False


def test_nothing_in_the_zendesk_tab_creates_a_new_article_draft(conn, qapp):
    """zendesk.md ("If it doesn't"): "There is no way to start a new article
    draft. This is a genuine gap in the current build — nothing in this tab
    creates one, and the drafts you see in the demo data were seeded rather
    than authored.\""""
    from src.ui.pages.enablement import zendesk_tab as zt
    src = Path(zt.__file__).read_text(encoding="utf-8")
    assert "save_article_draft" not in src
    assert "new_article" not in src
    # the only in-app caller of save_article_draft is the demo seeder
    hits = []
    for py in (_PROJECT_ROOT / "src").rglob("*.py"):
        if "save_article_draft" in py.read_text(encoding="utf-8"):
            hits.append(py.name)
    assert sorted(hits) == ["enablement_sim.py", "zendesk_store.py"]


def test_syncing_without_credentials_says_to_connect_zendesk_first(conn):
    """zendesk.md ("If it doesn't"): "Syncing says to connect Zendesk first.
    No credentials are configured.\""""
    page = _FakePage(conn)
    page._zd_client = None
    _page_cls()._on_zendesk_sync(page)
    assert page.zendesk.status == ["Connect Zendesk first (Settings)."]


def test_a_zendesk_api_error_surfaces_as_a_push_failure(conn):
    """zendesk.md ("If it doesn't"): "Pushing fails with a Zendesk error. The
    API rejected it.\""""
    class _Angry(_FakeZendeskClient):
        def update_article(self, *a, **k):
            raise RuntimeError("422 Unprocessable Entity")

    draft_id = zendesk_store.save_article_draft(conn, title="T", body="b",
                                                article_id=1)
    res = zendesk_store.publish_article_draft(conn, draft_id,
                                              zendesk_client=_Angry())
    assert res["ok"] is False
    assert res["error"].startswith("zendesk_push_failed")
    assert zendesk_store.get_article_draft(conn, draft_id)["status"] == "pending"


# ══════════════════════════════════════════════════════════════════════
#  article: zendesk-macros.md — "Editing macros"
# ══════════════════════════════════════════════════════════════════════

def _macro_draft(conn, actions, *, name="Refund reply", macro_id=None):
    return zendesk_store.save_macro_draft(conn, name=name, actions=actions,
                                          macro_id=macro_id,
                                          description="do not lose me")


def test_selecting_a_macro_shows_its_name_and_first_comment_action(conn, qapp):
    """zendesk-macros.md: "Selecting a macro draft fills in its name and puts
    its public reply in a plain-text box. The reply shown is the **first**
    comment action ...; if a macro somehow carries more than one, you only ever
    see the first.\""""
    from src.ui.pages.enablement.zendesk_tab import ZendeskPage
    draft_id = _macro_draft(conn, [
        {"field": "status", "value": "solved"},
        {"field": "comment_value", "value": "FIRST reply"},
        {"field": "comment_value", "value": "SECOND reply"},
    ])
    tab = ZendeskPage()
    page = _FakePage(conn)
    page.zendesk = tab
    _page_cls()._on_zd_macro_selected(page, draft_id)

    assert tab._m_name.text() == "Refund reply"
    assert tab._m_reply.toPlainText() == "FIRST reply"


def test_saving_a_macro_replaces_the_name_and_the_reply(conn):
    """zendesk-macros.md table: "Name → Replaced with the name field" and
    "Public reply → Replaced with the box, as plain text.\""""
    draft_id = _macro_draft(conn, [{"field": "comment_value", "value": "old"}])
    page = _FakePage(conn)
    _page_cls()._on_zd_macro_saved(page, draft_id, "New name", "new reply")
    d = zendesk_store.get_macro_draft(conn, draft_id)
    assert d["name"] == "New name"
    assert d["actions"][0] == {"field": "comment_value", "value": "new reply"}


def test_saving_a_macro_stores_the_reply_as_plain_text_losing_rich_html(conn):
    """zendesk-macros.md table: "Formatting in the reply → Lost — stored as a
    plain-text comment.\""""
    draft_id = _macro_draft(conn, [
        {"field": "comment_value_html", "value": "<b>Bold</b> <a href='#'>link</a>"},
    ])
    page = _FakePage(conn)
    _page_cls()._on_zd_macro_saved(page, draft_id, "Refund reply",
                                   "Bold link")
    actions = zendesk_store.get_macro_draft(conn, draft_id)["actions"]
    assert [a["field"] for a in actions] == ["comment_value"]
    assert "<b>" not in json.dumps(actions)


def test_saving_a_macro_collapses_a_second_comment_action(conn):
    """zendesk-macros.md table: "A second comment action → Collapsed into the
    one reply.\""""
    draft_id = _macro_draft(conn, [
        {"field": "comment_value", "value": "one"},
        {"field": "comment_value", "value": "two"},
        {"field": "comment_value_html", "value": "<p>three</p>"},
    ])
    page = _FakePage(conn)
    _page_cls()._on_zd_macro_saved(page, draft_id, "Refund reply", "merged")
    actions = zendesk_store.get_macro_draft(conn, draft_id)["actions"]
    assert actions == [{"field": "comment_value", "value": "merged"}]


def test_saving_a_macro_keeps_every_other_action_and_moves_the_reply_first(conn):
    """zendesk-macros.md table: "Tags, status, assignee, other actions → Kept",
    "Order of the actions → Reply moves to the front", and "A save should leave
    every action you did not edit intact.\""""
    draft_id = _macro_draft(conn, [
        {"field": "set_tags", "value": "refund"},
        {"field": "status", "value": "solved"},
        {"field": "comment_value", "value": "old"},
        {"field": "assignee_id", "value": "1234"},
    ])
    page = _FakePage(conn)
    _page_cls()._on_zd_macro_saved(page, draft_id, "Refund reply", "new")
    actions = zendesk_store.get_macro_draft(conn, draft_id)["actions"]
    assert actions[0] == {"field": "comment_value", "value": "new"}
    assert actions[1:] == [
        {"field": "set_tags", "value": "refund"},
        {"field": "status", "value": "solved"},
        {"field": "assignee_id", "value": "1234"},
    ]


def test_saving_a_macro_keeps_the_description_which_is_not_editable(conn, qapp):
    """zendesk-macros.md table: "Description → Kept, and not editable here.\""""
    from src.ui.pages.enablement import zendesk_tab as zt
    draft_id = _macro_draft(conn, [{"field": "comment_value", "value": "x"}])
    page = _FakePage(conn)
    _page_cls()._on_zd_macro_saved(page, draft_id, "Refund reply", "y")
    assert zendesk_store.get_macro_draft(conn, draft_id)["description"] == \
        "do not lose me"
    # the tab never references a description, so it cannot edit one
    assert "description" not in Path(zt.__file__).read_text(
        encoding="utf-8").lower()


def test_saving_a_macro_never_contacts_zendesk(conn):
    """zendesk-macros.md: "Nothing should reach Zendesk until you push. Editing
    and saving are local.\""""
    draft_id = _macro_draft(conn, [{"field": "comment_value", "value": "x"}])
    page = _FakePage(conn)
    client = _FakeZendeskClient()
    page._zd_client = client
    _page_cls()._on_zd_macro_saved(page, draft_id, "N", "reply")
    assert client.calls == []


def test_pushing_a_linked_macro_replaces_its_whole_action_list(conn):
    """zendesk-macros.md: "A draft linked to a live macro **replaces** that
    macro's actions with the local list rather than merging into it.\""""
    client = _FakeZendeskClient()
    actions = [{"field": "comment_value", "value": "hi"},
               {"field": "set_tags", "value": "refund"}]
    draft_id = _macro_draft(conn, actions, macro_id=99)
    zendesk_store.publish_macro_draft(conn, draft_id, zendesk_client=client)
    call = [c for c in client.calls if c[0] == "update_macro"][0]
    assert call[1] == 99
    assert call[3] == actions        # the full local list, sent verbatim


def test_pushing_an_unlinked_macro_creates_a_new_one(conn):
    """zendesk-macros.md: "An unlinked draft creates a new macro instead.\""""
    client = _FakeZendeskClient()
    draft_id = _macro_draft(conn, [{"field": "comment_value", "value": "hi"}])
    res = zendesk_store.publish_macro_draft(conn, draft_id,
                                            zendesk_client=client)
    assert [c[0] for c in client.calls] == ["create_macro"]
    assert res["macro_id"] == 777


def test_pushing_a_macro_with_no_connection_marks_it_pushed_locally(conn):
    """zendesk-macros.md: "pushing while Zendesk is not connected marks the
    draft pushed locally without contacting the API.\""""
    draft_id = _macro_draft(conn, [{"field": "comment_value", "value": "hi"}])
    res = zendesk_store.publish_macro_draft(conn, draft_id,
                                            zendesk_client=None)
    assert res["ok"] is True
    assert zendesk_store.get_macro_draft(conn, draft_id)["status"] == "pushed"
    assert zendesk_store.list_macro_drafts(conn) == []


def test_the_macro_editor_covers_only_the_name_and_the_reply(qapp):
    """zendesk-macros.md ("If it doesn't"): "You need to change something other
    than the reply. Not possible here; this editor covers the name and the
    reply only.\""""
    from src.ui.pages.enablement.zendesk_tab import ZendeskPage
    tab = ZendeskPage()
    saved = []
    tab.macro_saved.connect(lambda *a: saved.append(a))
    tab.show_macro_draft({"id": 3, "name": "M",
                          "actions": [{"field": "comment_value", "value": "r"},
                                      {"field": "set_tags", "value": "t"}]})
    tab._on_macro_save()
    assert saved == [(3, "M", "r")]      # only (draft_id, name, reply)


# ══════════════════════════════════════════════════════════════════════
#  article: content-studio.md — "The Content Studio"
# ══════════════════════════════════════════════════════════════════════

def _tools():
    from src.data.chat_tools.registry import get_tool_registry
    return get_tool_registry()


@pytest.fixture
def studio_env(conn, tmp_path, monkeypatch):
    """Isolate artifact files from the real project data/ dir and give the
    session-scoped tools a session id."""
    monkeypatch.setattr(artifact_store, "_ARTIFACTS_ROOT", tmp_path / "artifacts")
    sid_file = tmp_path / "session.txt"
    sid_file.write_text("sess-test", encoding="utf-8")
    monkeypatch.setenv("ALMA_CHAT_SESSION_FILE", str(sid_file))
    return conn


def _stub_llm(monkeypatch, replies):
    """Route llm_gen at a scripted client; returns the call log."""
    from src.gemini import client_factory
    calls = []
    seq = list(replies)

    class _C:
        def generate(self, prompt):
            calls.append(prompt)
            return seq.pop(0) if seq else ""

    monkeypatch.setattr(client_factory, "build_client_for_task",
                        lambda *a, **k: _C())
    return calls


def _no_llm(monkeypatch):
    from src.gemini import client_factory
    monkeypatch.setattr(client_factory, "build_client_for_task",
                        lambda *a, **k: None)


_GOOD_MERMAID = "```mermaid\nflowchart TD\n  A[\"Start\"] --> B[\"End\"]\n```"
_GOOD_QUIZ = """---
title: Refunds
pass_threshold: 70
---
### Q1: How long is the refund window?
- [ ] 7 days
- [x] 30 days
Explanation: Policy v2.
### Q2: Who approves exceptions?
- [x] The lead
- [ ] Anyone
"""
_GOOD_ONE_PAGER = ("# Refunds\n## Overview\no\n## Why it matters\nw\n"
                   "## Key facts\nk\n## How to talk about it\nh\n## Links\nl\n")


def test_the_studio_registers_exactly_the_documented_generators():
    """content-studio.md: "Four things can be generated" (diagram, quiz,
    one-pager/battle card via one doc tool, deck) plus attach / list / upload.
    studio-no-screen.md: these are reachable only as Renn tools."""
    tools = _tools()
    for name in ("generate_diagram", "generate_quiz", "generate_doc",
                 "generate_deck", "attach_artifact_to_draft", "list_artifacts",
                 "request_upload_artifact_to_drive"):
        assert name in tools, f"{name} is not a registered chat tool"


def test_one_pager_and_battle_card_are_the_two_doc_styles(studio_env):
    """content-studio.md: the one-pager and the battle card are the two
    "structured summary" generators, each with "Fixed section headings"."""
    res = artifact_tools.handle_generate_doc(studio_env, {"style": "novel"}, {})
    assert res["ok"] is False and res["error"] == "unknown_style"
    assert res["allowed_styles"] == ["battle_card", "one_pager"]
    assert artifact_tools._DOC_STYLES["one_pager"]["headers"] == [
        "## Overview", "## Why it matters", "## Key facts",
        "## How to talk about it", "## Links"]
    assert artifact_tools._DOC_STYLES["battle_card"]["headers"] == [
        "## Positioning", "## Objections and responses", "## Proof points",
        "## Landmines"]


def test_every_generator_requires_exactly_one_source(studio_env):
    """content-studio.md: "**Every generator takes exactly one source.** ...
    Not two, not zero. If you name several, Renn is told to come back and ask
    which one you meant rather than guessing.\""""
    handlers = (artifact_tools.handle_generate_diagram,
                artifact_tools.handle_generate_quiz,
                artifact_tools.handle_generate_deck)
    for handler in handlers:
        none = handler(studio_env, {}, {})
        assert none["error"] == "source_required", handler.__name__
        two = handler(studio_env, {"source_text": "a", "source_doc_id": "d"}, {})
        assert two["error"] == "source_required", handler.__name__
        assert "exactly ONE" in two["message"]
    doc_two = artifact_tools.handle_generate_doc(
        studio_env, {"style": "one_pager", "source_text": "a",
                     "source_task_id": "t"}, {})
    assert doc_two["error"] == "source_required"


def test_the_four_accepted_source_kinds_are_task_research_doc_and_text():
    """content-studio.md: "A task, a stored document, or text you paste in ...
    Research is also offered as a source type"."""
    assert artifact_tools._SOURCE_ARGS == (
        "source_task_id", "source_research_id", "source_doc_id", "source_text")


def test_generation_validates_in_code_and_repairs_exactly_once(studio_env,
                                                               monkeypatch):
    """content-studio.md: "Output is checked by code, not by the model's own
    judgement, and one repair attempt is made before giving up." and "If the
    second attempt still fails, you get an error rather than a broken
    artifact.\""""
    calls = _stub_llm(monkeypatch, ["not a diagram", "still not a diagram"])
    res = artifact_tools.handle_generate_diagram(
        studio_env, {"source_text": "some process"}, {})
    assert len(calls) == 2, "exactly one repair attempt"
    assert "failed validation" in calls[1]      # errors fed back verbatim
    assert res["ok"] is False
    assert res["error"] == "validation_failed"
    assert artifact_store.list_artifacts(studio_env) == []   # no half artifact


def test_a_repaired_second_attempt_succeeds_and_is_stored(studio_env,
                                                          monkeypatch):
    """content-studio.md: "one repair attempt is made before giving up" — the
    repair actually lands when the model corrects itself."""
    _stub_llm(monkeypatch, ["garbage", _GOOD_MERMAID])
    res = artifact_tools.handle_generate_diagram(
        studio_env, {"source_text": "steps"}, {})
    assert res["ok"] is True
    assert res["lint_retries"] == 1
    stored = artifact_store.get_artifact(studio_env, res["artifact_id"])
    assert stored["kind"] == "diagram"
    assert "flowchart TD" in json.loads(stored["spec_json"])["mermaid"]


def test_a_diagram_must_lint_as_valid_mermaid():
    """content-studio.md: "A diagram must lint as valid Mermaid"."""
    ok, errors, _ = mermaid_lint.lint("flowchart TD\n  A[\"x\"] --> B[\"y\"]")
    assert ok is True and errors == []
    bad_header, errs, _ = mermaid_lint.lint("here is your diagram:\nA --> B")
    assert bad_header is False
    assert any("missing diagram header" in e for e in errs)


def test_embedded_script_in_a_diagram_is_rejected_outright():
    """content-studio.md: "anything that looks like embedded script is rejected
    outright.\""""
    ok, errors, _ = mermaid_lint.lint(
        'flowchart TD\n  A["x"] --> B["y"]\n  click A "javascript:alert(1)"')
    assert ok is False
    assert any("javascript" in e for e in errors)
    ok2, errors2, _ = mermaid_lint.lint(
        'flowchart TD\n  A["x"] --> B["y"]\n  click A callback')
    assert ok2 is False
    assert any("callback" in e for e in errors2)


def test_a_quiz_needs_two_choices_and_a_correct_answer_per_question():
    """content-studio.md: "A quiz must have at least two choices and at least
    one correct answer per question.\""""
    one_choice = "### Q1: only one?\n- [x] yes\n"
    ok, errors, _ = quiz_artifacts.parse_quiz_md(one_choice)
    assert ok is False
    assert any("at least 2 choices" in e for e in errors)

    none_correct = "### Q1: which?\n- [ ] a\n- [ ] b\n"
    ok2, errors2, _ = quiz_artifacts.parse_quiz_md(none_correct)
    assert ok2 is False
    assert any("no correct answer" in e for e in errors2)


def test_a_quiz_is_clamped_to_between_two_and_twelve_questions(studio_env,
                                                               monkeypatch):
    """content-studio.md table: "Quiz → A knowledge check → Two to twelve
    questions". The requested count is clamped into [2, 12] before it reaches
    the prompt, and more than twelve fails validation outright."""
    template = (_PROJECT_ROOT / "config" / "prompts"
                / "enablement_quiz.txt").read_text(encoding="utf-8")
    head, tail = template.split("{n_questions}", 1)

    def _requested(prompt):
        """Recover the number the handler substituted into the template."""
        rest = prompt[len(head):]
        return int(re.match(r"\d+", rest).group(0))

    prompts = _stub_llm(monkeypatch, [_GOOD_QUIZ, _GOOD_QUIZ, _GOOD_QUIZ])
    for asked, expected in ((1, 2), (99, 12), (5, 5)):
        prompts.clear()
        artifact_tools.handle_generate_quiz(
            studio_env, {"source_text": "s", "n_questions": asked}, {})
        assert prompts, f"no prompt issued for n_questions={asked}"
        assert _requested(prompts[0]) == expected

    assert quiz_artifacts.MAX_QUESTIONS == 12
    ok, errors, _ = quiz_artifacts.parse_quiz_md(
        "\n".join(f"### Q{i}: q\n- [x] a\n- [ ] b" for i in range(1, 14)))
    assert ok is False and any("too many questions" in e for e in errors)


def test_a_one_pager_must_contain_all_its_required_headings(studio_env,
                                                            monkeypatch):
    """content-studio.md: "A one-pager or battle card must contain all of its
    required section headings.\""""
    missing = "# Refunds\n## Overview\nonly this one\n"
    _stub_llm(monkeypatch, [missing, missing])
    res = artifact_tools.handle_generate_doc(
        studio_env, {"style": "one_pager", "source_text": "s"}, {})
    assert res["ok"] is False and res["error"] == "validation_failed"
    assert any("Why it matters" in e for e in res["format_errors"])
    assert artifact_store.list_artifacts(studio_env) == []

    _stub_llm(monkeypatch, [_GOOD_ONE_PAGER])
    ok = artifact_tools.handle_generate_doc(
        studio_env, {"style": "one_pager", "source_text": "s"}, {})
    assert ok["ok"] is True
    assert artifact_store.get_artifact(studio_env,
                                       ok["artifact_id"])["kind"] == "one_pager"


def test_no_model_means_nothing_is_created(studio_env, monkeypatch):
    """content-studio.md ("If it doesn't"): "Generation fails saying no model is
    available. Nothing is created; there is no half-made artifact to clean
    up." (diagram / quiz / one-pager — decks are documented as the exception)"""
    _no_llm(monkeypatch)
    for res in (
        artifact_tools.handle_generate_diagram(studio_env,
                                               {"source_text": "s"}, {}),
        artifact_tools.handle_generate_quiz(studio_env,
                                            {"source_text": "s"}, {}),
        artifact_tools.handle_generate_doc(
            studio_env, {"style": "battle_card", "source_text": "s"}, {}),
    ):
        assert res == {"ok": False, "error": "no_llm_client"}
    assert artifact_store.list_artifacts(studio_env) == []


def test_a_deck_falls_back_to_a_mechanical_outline_instead_of_failing(
        studio_env, monkeypatch):
    """content-studio.md: "if the model cannot produce a usable outline, the
    deck falls back to a mechanical heading-based outline rather than failing,
    so you still get a file.\""""
    _no_llm(monkeypatch)
    source = "# Refund windows\n- 30 days\n\n# Exceptions\n- Lead approval\n"
    res = artifact_tools.handle_generate_deck(
        studio_env, {"source_text": source, "title": "Refunds"}, {})
    assert res["ok"] is True, res
    assert Path(res["file_path"]).exists()
    art = artifact_store.get_artifact(studio_env, res["artifact_id"])
    assert art["kind"] == "deck" and art["status"] == "rendered"
    assert json.loads(art["provenance_json"])["outline_via"] == "deterministic"
    deck = pptx_store.get_deck(studio_env, res["deck_id"])
    assert [s["title"] for s in deck["outline"]["slides"]][1:] == [
        "Refund windows", "Exceptions"]


def test_deck_generation_reports_progress_as_a_tracked_job(studio_env,
                                                           monkeypatch):
    """content-studio.md: "Decks ... run as a tracked job" / "Deck generation
    reports progress as a job you can watch.\""""
    from src.data import agent_jobs
    _no_llm(monkeypatch)
    res = artifact_tools.handle_generate_deck(
        studio_env, {"source_text": "# A\n- b\n"}, {})
    job = agent_jobs.get_job(studio_env, res["job_id"])
    assert job["kind"] == "deck_generation"
    assert job["status"] == "done"
    steps = [s["name"] for s in agent_jobs.list_steps(studio_env, res["job_id"])]
    assert steps == ["load source", "outline", "export"]


def test_attaching_appends_to_a_draft_marks_attached_and_does_not_publish(
        studio_env, monkeypatch):
    """content-studio.md: "**Anything you generate can be attached to a card
    draft.** Attaching appends the content to the draft and marks the artifact
    attached. It does not publish.\""""
    from src.data import enablement_store
    _stub_llm(monkeypatch, [_GOOD_ONE_PAGER])
    made = artifact_tools.handle_generate_doc(
        studio_env, {"style": "one_pager", "source_text": "s"}, {})
    draft_id = enablement_store.save_card_draft(studio_env, title="Card",
                                                content="EXISTING BODY")
    res = artifact_tools.handle_attach_artifact(
        studio_env, {"artifact_id": made["artifact_id"],
                     "draft_id": draft_id}, {})
    assert res["ok"] is True
    draft = enablement_store.get_draft(studio_env, draft_id)
    assert draft["content"].startswith("EXISTING BODY")      # appended
    assert "## Why it matters" in draft["content"]
    assert draft["status"] == "pending"                      # not published
    art = artifact_store.get_artifact(studio_env, made["artifact_id"])
    assert art["status"] == "attached"
    assert art["draft_id"] == draft_id


def test_a_deck_cannot_be_attached_to_a_card(studio_env, monkeypatch):
    """content-studio.md: "Decks are the exception — they do not attach to
    cards". studio-no-screen.md: "Renn says a deck cannot be attached to a card.
    Correct.\""""
    _no_llm(monkeypatch)
    made = artifact_tools.handle_generate_deck(
        studio_env, {"source_text": "# A\n- b\n"}, {})
    res = artifact_tools.handle_attach_artifact(
        studio_env, {"artifact_id": made["artifact_id"],
                     "new_draft_title": "Deck card"}, {})
    assert res["ok"] is False
    assert res["error"] == "kind_not_attachable"


def test_uploading_a_deck_opens_a_confirm_card_naming_the_file(studio_env,
                                                               monkeypatch):
    """content-studio.md: "Uploading a deck should always show a Confirm card
    naming the file"."""
    _no_llm(monkeypatch)
    made = artifact_tools.handle_generate_deck(
        studio_env, {"source_text": "# A\n- b\n", "title": "Refund Deck"}, {})
    out = artifact_tools.handle_request_upload_artifact(
        studio_env, {"artifact_id": made["artifact_id"],
                     "target_folder_id": "FOLDER-123"}, {})
    assert isinstance(out, str) and "confirmation card" in out.lower()
    row = studio_env.execute(
        "SELECT payload_json FROM chat_action_requests WHERE type='confirm_write'"
    ).fetchone()
    payload = json.loads(row[0])
    assert payload["op"] == "upload_artifact_to_drive"
    assert "refund-deck.pptx" in payload["summary"]
    assert payload["params"]["folder_id"] == "FOLDER-123"


def test_upload_targets_the_named_folder_or_the_kb_folder_never_the_active_one(
        studio_env, monkeypatch):
    """content-studio.md: "it should target the knowledge-base folder or a
    folder you name — never whichever folder happens to be active.\""""
    from src.data import settings_manager
    _no_llm(monkeypatch)
    made = artifact_tools.handle_generate_deck(
        studio_env, {"source_text": "# A\n- b\n"}, {})

    # no explicit folder → falls back to enablement.kb.ec_folder_id only
    monkeypatch.setattr(settings_manager, "get_section",
                        lambda name, default=None: {
                            "kb": {"ec_folder_id": "EC-FOLDER"},
                            "active_folders": ["PHI-PRODUCT-FOLDER"],
                        } if name == "enablement" else (default or {}))
    artifact_tools.handle_request_upload_artifact(
        studio_env, {"artifact_id": made["artifact_id"]}, {})
    payloads = [json.loads(r[0]) for r in studio_env.execute(
        "SELECT payload_json FROM chat_action_requests").fetchall()]
    assert payloads[-1]["params"]["folder_id"] == "EC-FOLDER"
    assert "PHI-PRODUCT-FOLDER" not in json.dumps(payloads)


def test_upload_without_a_kb_folder_steers_instead_of_guessing(studio_env,
                                                               monkeypatch):
    """content-studio.md ("If it doesn't"): "Uploading says no knowledge-base
    folder is set up.\""""
    from src.data import settings_manager
    _no_llm(monkeypatch)
    made = artifact_tools.handle_generate_deck(
        studio_env, {"source_text": "# A\n- b\n"}, {})
    monkeypatch.setattr(settings_manager, "get_section",
                        lambda name, default=None: default or {})
    res = artifact_tools.handle_request_upload_artifact(
        studio_env, {"artifact_id": made["artifact_id"]}, {})
    assert res["ok"] is False and res["error"] == "ec_not_bootstrapped"


def test_uploading_an_unrendered_artifact_says_it_has_no_file(studio_env,
                                                              monkeypatch):
    """content-studio.md ("If it doesn't"): "Uploading says the artifact has not
    been rendered. Only decks produce a file today.\""""
    _stub_llm(monkeypatch, [_GOOD_MERMAID])
    made = artifact_tools.handle_generate_diagram(
        studio_env, {"source_text": "s"}, {})
    assert artifact_store.get_artifact(
        studio_env, made["artifact_id"])["file_path"] is None
    res = artifact_tools.handle_request_upload_artifact(
        studio_env, {"artifact_id": made["artifact_id"]}, {})
    assert res["ok"] is False and res["error"] == "artifact_not_rendered"


def test_every_artifact_records_what_it_was_built_from(studio_env, monkeypatch):
    """content-studio.md: "Generated content should be traceable to its source.
    Each artifact records what it was built from.\""""
    _stub_llm(monkeypatch, [_GOOD_MERMAID, _GOOD_QUIZ, _GOOD_ONE_PAGER])
    ids = [
        artifact_tools.handle_generate_diagram(
            studio_env, {"source_text": "s"}, {})["artifact_id"],
        artifact_tools.handle_generate_quiz(
            studio_env, {"source_text": "s"}, {})["artifact_id"],
        artifact_tools.handle_generate_doc(
            studio_env, {"style": "one_pager", "source_text": "s"},
            {})["artifact_id"],
    ]
    for aid in ids:
        prov = json.loads(artifact_store.get_artifact(
            studio_env, aid)["provenance_json"])
        assert prov["source_ref"] == "inline"


def test_a_document_source_is_recorded_as_a_link_and_a_source_ref(studio_env,
                                                                  monkeypatch):
    """content-studio.md: "Each artifact records what it was built from" — a
    stored document source is linked by id, not just described."""
    from src.data import enablement_store
    doc_id = enablement_store.save_document(
        studio_env, doc_id="doc-1", name="Refund policy",
        full_text="# Refunds\nThirty days.", source="local")
    _stub_llm(monkeypatch, [_GOOD_MERMAID])
    res = artifact_tools.handle_generate_diagram(
        studio_env, {"source_doc_id": "doc-1"}, {})
    art = artifact_store.get_artifact(studio_env, res["artifact_id"])
    assert art["doc_id"] == "doc-1"
    assert json.loads(art["provenance_json"])["source_ref"] == "doc:doc-1"
    assert doc_id


# ══════════════════════════════════════════════════════════════════════
#  article: studio-no-screen.md — "The Content Studio has no screen yet"
# ══════════════════════════════════════════════════════════════════════

def _enablement_tab_map() -> dict:
    """Extract the literal ``self._tab_widgets`` mapping from page.py without
    constructing the page (which starts threads and reads live settings)."""
    tree = ast.parse(_PAGE_PY.read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        if (isinstance(node, ast.Assign)
                and isinstance(node.targets[0], ast.Attribute)
                and node.targets[0].attr == "_tab_widgets"
                and isinstance(node.value, ast.Dict)):
            return {k.value: True for k in node.value.keys
                    if isinstance(k, ast.Constant)}
    raise AssertionError("_tab_widgets literal not found in page.py")


def _enablement_tab_labels() -> list[str]:
    tree = ast.parse(_PAGE_PY.read_text(encoding="utf-8"))
    labels = []
    for node in ast.walk(tree):
        if (isinstance(node, ast.Call)
                and isinstance(node.func, ast.Attribute)
                and node.func.attr == "addTab"
                and isinstance(node.func.value, ast.Attribute)
                and node.func.value.attr == "tabs"
                and len(node.args) == 2
                and isinstance(node.args[1], ast.Constant)):
            labels.append(node.args[1].value)
    return labels


def test_there_is_no_content_studio_tab_menu_entry_or_button():
    """studio-no-screen.md: "There is no studio tab, no menu entry, no button.
    Asking Renn is the only way in.\""""
    keys = _enablement_tab_map()
    labels = _enablement_tab_labels()
    assert set(keys) == {"home", "calendar", "tasks", "workbench", "analytics",
                         "powerpoint", "zendesk", "help", "settings"}
    assert labels == ["Home", "Calendar", "Tasks", "Workbench", "Analytics",
                      "PowerPoint", "Zendesk", "Help", "Settings"]
    joined = " ".join(list(keys) + labels).lower()
    assert "studio" not in joined


def test_no_ui_module_calls_a_studio_generator():
    """studio-no-screen.md: "nothing in the interface points at them" — the
    generators have no caller under src/ui/."""
    generators = ("handle_generate_diagram", "handle_generate_quiz",
                  "handle_generate_doc", "handle_generate_deck")
    offenders = []
    for py in (_PROJECT_ROOT / "src" / "ui").rglob("*.py"):
        text = py.read_text(encoding="utf-8", errors="ignore")
        if any(g in text for g in generators):
            offenders.append(str(py.relative_to(_PROJECT_ROOT)))
    assert offenders == []


def test_listing_artifacts_enumerates_them_without_their_contents(studio_env,
                                                                  monkeypatch):
    """studio-no-screen.md: "ask for a list ... This enumerates them; it does
    not show their contents.\""""
    _stub_llm(monkeypatch, [_GOOD_MERMAID, _GOOD_QUIZ])
    artifact_tools.handle_generate_diagram(studio_env, {"source_text": "s"}, {})
    artifact_tools.handle_generate_quiz(studio_env, {"source_text": "s"}, {})
    res = artifact_tools.handle_list_artifacts(studio_env, {}, {})
    assert res["count"] == 2
    for row in res["artifacts"]:
        assert "spec_json" not in row and "quiz_md" not in json.dumps(row)
        assert set(row) == {"artifact_id", "kind", "title", "status", "task_id",
                            "card_id", "draft_id", "file_path", "updated_at"}
    filtered = artifact_tools.handle_list_artifacts(studio_env,
                                                    {"kind": "quiz"}, {})
    assert filtered["count"] == 1


def test_a_renn_generated_deck_shows_up_in_the_powerpoint_tab_list(studio_env,
                                                                   monkeypatch,
                                                                   qapp):
    """studio-no-screen.md: "**Decks are the one kind you can see in the
    interface.** A deck Renn generates is saved as a real deck and appears in
    the PowerPoint tab's list alongside decks you modelled yourself.\""""
    from src.ui.pages.enablement.pptx_tab import PptxPage
    _no_llm(monkeypatch)
    hand_made = _mk_deck(studio_env, "Modelled by hand",
                         [{"title": "A", "bullets": ["x"]}])
    res = artifact_tools.handle_generate_deck(
        studio_env, {"source_text": "# Renn\n- built this\n",
                     "title": "Built by Renn"}, {})
    tab = PptxPage()
    page = _FakePage(studio_env)
    page.pptx = tab
    _page_cls()._load_pptx(page)

    titles = [tab._list.item(i).text() for i in range(tab._list.count())]
    assert any("Built by Renn" in t for t in titles)
    assert any("Modelled by hand" in t for t in titles)
    assert res["deck_id"] != hand_made


def test_only_the_five_documented_kinds_are_generatable():
    """studio-no-screen.md: "Only diagrams, quizzes, one-pagers, battle cards
    and decks exist. Anything else is invented"."""
    generatable = {"diagram", "quiz", "one_pager", "battle_card", "deck"}
    assert artifact_store.KINDS == generatable | {"podcast"}
    assert artifact_store.KINDS - generatable == {"podcast"}


# ══════════════════════════════════════════════════════════════════════
#  article: deferred-previews.md
# ══════════════════════════════════════════════════════════════════════

def test_generating_a_diagram_stores_source_and_mints_no_render(studio_env,
                                                                monkeypatch):
    """deferred-previews.md: "Generating a diagram produces valid Mermaid source
    and stores it. No picture is drawn.\""""
    _stub_llm(monkeypatch, [_GOOD_MERMAID])
    res = artifact_tools.handle_generate_diagram(
        studio_env, {"source_text": "s"}, {})
    art = artifact_store.get_artifact(studio_env, res["artifact_id"])
    assert art["status"] == "draft"
    assert art["file_path"] is None
    assert json.loads(art["spec_json"])["mermaid"].startswith("flowchart TD")
    assert "renders when previews ship" in res["note"]


def test_there_is_no_mermaid_renderer_anywhere_in_the_app():
    """deferred-previews.md: "There is no viewer, no thumbnail, and no render
    step anywhere in the app.\""""
    pkg = json.loads((_PROJECT_ROOT / "web" / "package.json").read_text(
        encoding="utf-8"))
    deps = {**pkg.get("dependencies", {}), **pkg.get("devDependencies", {})}
    assert not any("mermaid" in d.lower() for d in deps)

    # no mermaid runtime is imported or invoked anywhere in the web layer
    js_markers = re.compile(
        r"mermaid\.(render|init|initialize|parse)|from\s+['\"]mermaid['\"]"
        r"|import\s+mermaid|require\(['\"]mermaid['\"]\)")
    offenders = []
    for pattern in ("*.js", "*.jsx", "*.ts", "*.tsx"):
        for f in (_PROJECT_ROOT / "web" / "src").rglob(pattern):
            if js_markers.search(f.read_text(encoding="utf-8", errors="ignore")):
                offenders.append(str(f.relative_to(_PROJECT_ROOT)))
    assert offenders == []

    # no Python module imports a mermaid rendering runtime either
    py_import = re.compile(r"^\s*(?:import\s+mermaid|from\s+mermaid\b)",
                           re.MULTILINE)
    py_offenders = [
        f.relative_to(_PROJECT_ROOT).as_posix()
        for f in (_PROJECT_ROOT / "src").rglob("*.py")
        if py_import.search(f.read_text(encoding="utf-8", errors="ignore"))]
    assert py_offenders == []

    # the one mermaid-aware Python module is the linter, and it only validates:
    # its whole public surface is ``lint(text) -> (ok, errors, cleaned)``
    assert sorted(n for n in vars(mermaid_lint)
                  if not n.startswith("_")
                  and callable(getattr(mermaid_lint, n))) == ["lint"]

    # the reserved preview request type is declared but nothing emits it
    from src.data import chat_action_requests
    emitters = [
        f.relative_to(_PROJECT_ROOT).as_posix()
        for f in (_PROJECT_ROOT / "src").rglob("*.py")
        if "artifact_preview" in f.read_text(encoding="utf-8", errors="ignore")]
    assert "artifact_preview" in chat_action_requests.INFORMATIONAL_ACTION_TYPES
    assert emitters == ["src/data/chat_action_requests.py"]


def test_attaching_a_diagram_inserts_its_source_in_a_fenced_block(studio_env,
                                                                  monkeypatch):
    """deferred-previews.md: "Attaching a diagram to a card draft inserts the
    diagram's **source text** in a fenced block, so the card carries the
    definition rather than an image.\""""
    from src.data import enablement_store
    _stub_llm(monkeypatch, [_GOOD_MERMAID])
    made = artifact_tools.handle_generate_diagram(
        studio_env, {"source_text": "s", "title": "Refund flow"}, {})
    res = artifact_tools.handle_attach_artifact(
        studio_env, {"artifact_id": made["artifact_id"],
                     "new_draft_title": "Refunds card"}, {})
    draft = enablement_store.get_draft(studio_env, res["draft_id"])
    assert "```mermaid" in draft["content"]
    assert "flowchart TD" in draft["content"]
    assert not draft.get("content_html")     # no image, no rendered HTML


def test_the_quiz_itself_is_real_and_fully_parsed():
    """deferred-previews.md: "The quiz itself is real — questions, choices,
    correct answers and explanations are generated and validated.\""""
    ok, errors, quiz = quiz_artifacts.parse_quiz_md(_GOOD_QUIZ)
    assert ok is True and errors == []
    assert quiz["title"] == "Refunds"
    assert quiz["pass_threshold"] == 70
    assert len(quiz["questions"]) == 2
    q1 = quiz["questions"][0]
    assert q1["question"] == "How long is the refund window?"
    assert [c["text"] for c in q1["choices"]] == ["7 days", "30 days"]
    assert [c["correct"] for c in q1["choices"]] == [False, True]
    assert q1["explanation"] == "Policy v2."


def test_there_is_no_interactive_quiz_card_in_the_app():
    """deferred-previews.md: "What is missing is the interactive card that would
    let you take the quiz inside the app." Nothing DEFINES a quiz component or
    widget (artifact_tools only mentions the deferred one in a docstring), and
    the quiz renderer that does exist emits no interactive markup."""
    definition = re.compile(
        r"^\s*(class\s+Quiz\w*|def\s+Quiz\w*|function\s+Quiz\w*"
        r"|(?:const|let|var)\s+Quiz\w*\s*=)", re.MULTILINE)
    offenders = []
    for root, patterns in ((_PROJECT_ROOT / "src", ("*.py",)),
                           (_PROJECT_ROOT / "web" / "src", ("*.js", "*.jsx"))):
        for pattern in patterns:
            for f in root.rglob(pattern):
                if definition.search(f.read_text(encoding="utf-8",
                                                 errors="ignore")):
                    offenders.append(str(f.relative_to(_PROJECT_ROOT)))
    assert offenders == []

    _, _, quiz = quiz_artifacts.parse_quiz_md(_GOOD_QUIZ)
    rendered = quiz_artifacts.to_guru_html(quiz).lower()
    for interactive in ("<script", "<input", "<button", "<form", "onclick"):
        assert interactive not in rendered


def test_attaching_a_quiz_produces_static_html_with_collapsible_answers(
        studio_env, monkeypatch):
    """deferred-previews.md: "Attaching a quiz to a card draft converts it into
    plain static HTML with each answer behind a collapsible section.\""""
    from src.data import enablement_store
    _stub_llm(monkeypatch, [_GOOD_QUIZ])
    made = artifact_tools.handle_generate_quiz(
        studio_env, {"source_text": "s"}, {})
    res = artifact_tools.handle_attach_artifact(
        studio_env, {"artifact_id": made["artifact_id"],
                     "new_draft_title": "Quiz card"}, {})
    html = enablement_store.get_draft(studio_env, res["draft_id"])["content_html"]
    assert html.count("<details><summary>Show answer</summary>") == 2
    assert "30 days" in html
    assert "<script" not in html.lower()
    assert "onclick" not in html.lower()


def test_podcast_is_a_reserved_kind_with_no_implementation():
    """deferred-previews.md: "**Podcasts do not exist.** The word appears in
    exactly one place: a reserved entry in the list of artifact kinds. There is
    no generator, no audio, no text-to-speech, no player.\""""
    assert "podcast" in artifact_store.KINDS

    # no generator tool, no dispatch, no audio/TTS anywhere
    tools = _tools()
    assert not [t for t in tools if "podcast" in t.lower()]
    audio = re.compile(r"text[_-]?to[_-]?speech|\btts\b|podcast", re.IGNORECASE)
    modules = []
    for f in (_PROJECT_ROOT / "src").rglob("*.py"):
        if audio.search(f.read_text(encoding="utf-8", errors="ignore")):
            modules.append(f.relative_to(_PROJECT_ROOT).as_posix())
    assert modules == ["src/data/artifact_store.py"], modules


def test_nothing_can_attach_or_upload_a_podcast(studio_env):
    """deferred-previews.md: "nothing that reads that entry ... Nothing can
    create one, and asking for one cannot work.\""""
    aid = artifact_store.create_artifact(studio_env, kind="podcast",
                                         title="Reserved seam")
    res = artifact_tools.handle_attach_artifact(
        studio_env, {"artifact_id": aid, "new_draft_title": "X"}, {})
    assert res["ok"] is False and res["error"] == "kind_not_attachable"
    up = artifact_tools.handle_request_upload_artifact(
        studio_env, {"artifact_id": aid}, {})
    assert up["ok"] is False and up["error"] == "artifact_not_rendered"


# ══════════════════════════════════════════════════════════════════════
#  corpus hygiene — the articles under audit exist and are loadable
# ══════════════════════════════════════════════════════════════════════

def test_all_seven_create_articles_are_present_and_parse():
    """Guards the audit itself: every article named in this file is on disk
    with the frontmatter fields the loader requires."""
    from src.data.help.loader import _parse_yaml, _split_frontmatter
    expected = {"content-studio.md", "deck-outline.md", "deferred-previews.md",
                "powerpoint.md", "studio-no-screen.md", "zendesk-macros.md",
                "zendesk.md"}
    assert {p.name for p in _HELP_DIR.glob("*.md")} == expected
    for path in sorted(_HELP_DIR.glob("*.md")):
        fm, body = _split_frontmatter(path.read_text(encoding="utf-8"))
        meta = _parse_yaml(fm)
        assert meta.get("section") == "create", path.name
        assert meta.get("id") and meta.get("title"), path.name
        assert "## How it works" in body, path.name
