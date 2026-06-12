"""E4 — pptx_store: outline CRUD, tolerant LLM parse, real .pptx export
(read back with python-pptx), demo seed. Headless (no Qt)."""

import json

import pytest

from src.data import pptx_store


class StubLLM:
    def __init__(self, payload):
        self.payload = payload
        self.prompts = []

    def generate(self, prompt):
        self.prompts.append(prompt)
        return self.payload


_OUTLINE = {
    "title": "SSO Rollout",
    "slides": [
        {"title": "What changed", "bullets": ["self-serve SSO", "June 24 rollout"]},
        {"title": "Steps", "bullets": ["open console", "choose IdP", "test pilot"]},
    ],
}


class TestCrud:
    def test_save_get_round_trip(self, empty_db):
        conn = empty_db.conn
        did = pptx_store.save_deck(conn, title="SSO Rollout", outline=_OUTLINE,
                                   source_ref="doc-1")
        deck = pptx_store.get_deck(conn, did)
        assert deck["title"] == "SSO Rollout"
        assert deck["slide_count"] == 2
        assert deck["outline"]["slides"][0]["title"] == "What changed"
        assert deck["status"] == "pending"

    def test_list_and_update(self, empty_db):
        conn = empty_db.conn
        did = pptx_store.save_deck(conn, title="D", outline={"title": "D", "slides": []})
        assert any(d["id"] == did for d in pptx_store.list_decks(conn))
        res = pptx_store.update_deck_outline(
            conn, did, title="D2",
            outline={"title": "D2", "slides": [{"title": "S", "bullets": ["b"]}]})
        assert res["ok"] and res["slides"] == 1
        assert pptx_store.get_deck(conn, did)["title"] == "D2"

    def test_normalizes_garbage_slides(self, empty_db):
        conn = empty_db.conn
        did = pptx_store.save_deck(conn, title="x", outline={
            "title": "x", "slides": ["not a dict", {"title": "ok", "bullets": [1, 2]}]})
        deck = pptx_store.get_deck(conn, did)
        assert len(deck["outline"]["slides"]) == 1
        assert deck["outline"]["slides"][0]["bullets"] == ["1", "2"]


class TestParseOutline:
    def test_raw_json(self):
        out = pptx_store.parse_outline(json.dumps(_OUTLINE))
        assert out["title"] == "SSO Rollout" and len(out["slides"]) == 2

    def test_json_in_prose_and_fences(self):
        wrapped = "Here is your deck:\n```json\n" + json.dumps(_OUTLINE) + "\n```\nEnjoy!"
        out = pptx_store.parse_outline(wrapped)
        assert out["title"] == "SSO Rollout"

    def test_garbage_returns_empty_deck(self):
        out = pptx_store.parse_outline("not json at all")
        assert out["title"] and out["slides"] == []


class TestGenerate:
    def test_from_document(self, empty_db):
        from src.data.enablement_store import save_document
        conn = empty_db.conn
        doc_id = save_document(conn, source="drive", name="SSO.gdoc",
                               full_text="SSO self-serve, June 24 rollout.")
        llm = StubLLM(json.dumps(_OUTLINE))
        res = pptx_store.generate_deck_from_document(conn, doc_id, llm)
        assert res["ok"] and res["slides"] == 2
        assert "SSO.gdoc" in llm.prompts[0]
        deck = pptx_store.get_deck(conn, res["deck_id"])
        assert deck["source_ref"] == doc_id

    def test_from_document_missing(self, empty_db):
        res = pptx_store.generate_deck_from_document(empty_db.conn, "nope", StubLLM("{}"))
        assert not res["ok"]


class TestExport:
    def test_real_pptx_written_and_readable(self, empty_db, tmp_path):
        conn = empty_db.conn
        did = pptx_store.save_deck(conn, title="SSO Rollout", outline=_OUTLINE)
        out = tmp_path / "deck.pptx"
        res = pptx_store.export_pptx(conn, did, str(out))
        assert res["ok"]
        assert out.exists() and out.stat().st_size > 0
        # read it back with python-pptx
        from pptx import Presentation
        prs = Presentation(str(out))
        # cover + 2 content slides
        assert len(prs.slides) == 3
        titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
        assert "SSO Rollout" in titles
        assert "What changed" in titles and "Steps" in titles
        # status flipped + file_path recorded
        deck = pptx_store.get_deck(conn, did)
        assert deck["status"] == "exported"
        assert deck["file_path"] == str(out)

    def test_export_missing_deck(self, empty_db, tmp_path):
        res = pptx_store.export_pptx(empty_db.conn, 999, str(tmp_path / "x.pptx"))
        assert not res["ok"]


class TestDemoSeed:
    def test_seed(self, empty_db):
        from src.data.enablement_sim import seed_demo_decks
        out = seed_demo_decks(empty_db.conn)
        assert len(out["decks"]) == 2
        decks = pptx_store.list_decks(empty_db.conn)
        assert len(decks) == 2
